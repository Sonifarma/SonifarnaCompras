"""Functional regressions against a disposable database, never the buyer's data."""
import io
import os
from pathlib import Path
import tempfile
import unittest
from openpyxl import Workbook, load_workbook
from fastapi.testclient import TestClient

sandbox = tempfile.TemporaryDirectory()
os.environ['SONIFARMA_COMPRAS_DB'] = str(Path(sandbox.name) / 'test.sqlite3')
import app as compras


class PurchaseWorkflow(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from auth import password_hash
        with compras.db() as c:
            c.execute('INSERT OR REPLACE INTO app_users(username,password_hash,must_change) VALUES (?,?,0)', ('workflow', password_hash('workflow-test-password')))
        cls.client = TestClient(compras.app, headers={'X-Sonifarma-Request': '1'})
        cls.client.post('/api/auth/login', data={'username': 'workflow', 'password': 'workflow-test-password'}).raise_for_status()

    def setUp(self):
        self.pid = self.client.post('/api/processes', data={'name': 'Teste isolado'}).json()['id']
        compras.add_need(self.pid, ['T1', 'Produto teste', '', 10, 'G', 5])
        compras.add_quote(self.pid, 'Fornecedor teste', 'Produto teste', '10 G', 10, 'G', 20, 2, '', 'teste')
        with compras.db() as c:
            self.need = c.execute('SELECT id FROM needs WHERE process_id=?', (self.pid,)).fetchone()[0]
            self.quote = c.execute('SELECT id FROM quotes WHERE process_id=?', (self.pid,)).fetchone()[0]

    def data(self):
        result = self.client.get(f'/api/data/{self.pid}')
        self.assertEqual(result.status_code, 200)
        return result.json()

    def test_removal_survives_refresh(self):
        self.assertEqual(self.data()['summary']['selected'], 1)
        self.client.post('/api/unchoose', data={'need_id': self.need}).raise_for_status()
        for _ in range(2):
            self.assertEqual(self.data()['summary']['selected'], 0)
        self.client.post('/api/choose', data={'need_id': self.need, 'quote_id': self.quote}).raise_for_status()
        self.assertEqual(self.data()['summary']['selected'], 1)

    def test_estimates_and_exports(self):
        self.assertEqual(self.data()['summary']['selected_total'], 20)
        self.assertEqual(self.data()['summary']['indicative_saving'], 30)
        xlsx = self.client.get(f'/api/drafts/{self.pid}/xlsx')
        self.assertEqual(xlsx.status_code, 200)
        workbook = load_workbook(io.BytesIO(xlsx.content))
        self.assertEqual(workbook.active['I5'].value, 20)
        for path in [f'/api/export/{self.pid}', f'/api/drafts/{self.pid}/pdf']:
            response = self.client.get(path)
            self.assertEqual(response.status_code, 200)
            self.assertTrue(response.content.startswith(b'%PDF'))

    def test_manual_choice_preserved(self):
        compras.add_quote(self.pid, 'Alternativo', 'Produto teste', '10 G', 10, 'G', 10, 1, '', 'teste')
        self.client.post('/api/choose', data={'need_id': self.need, 'quote_id': self.quote}).raise_for_status()
        self.assertEqual(self.data()['needs'][0]['choice']['quote_id'], self.quote)

    def test_cross_process_choice_rejected(self):
        other = compras.create_process('Outro')
        compras.add_need(other, ['X', 'Outro produto', '', 1, 'G', 1])
        with compras.db() as c:
            need = c.execute('SELECT id FROM needs WHERE process_id=?', (other,)).fetchone()[0]
        for path in ['/api/choose', '/api/match']:
            self.assertEqual(self.client.post(path, data={'need_id': need, 'quote_id': self.quote}).status_code, 400)

    def test_invalid_inputs(self):
        self.assertEqual(self.client.post('/api/processes', data={'name': '   '}).status_code, 400)
        self.assertEqual(self.client.get('/api/data/9999999').status_code, 404)
        for path, filename in [('needs', 'broken.xlsx'), ('quotes', 'broken.xlsx'), ('quotes', 'broken.pdf')]:
            response = self.client.post(f'/api/{path}/{self.pid}', data={'supplier': 'Teste'}, files={'file': (filename, b'corrupt')})
            self.assertEqual(response.status_code, 400)
        self.assertIsNone(compras.num('NaN'))
        self.assertIsNone(compras.num(float('inf')))

    def test_import_and_audit(self):
        workbook = Workbook()
        workbook.active.append(['Código', 'Produto', 'Fornecedor', 'Quantidade', 'Unidade', 'Custo'])
        workbook.active.append(['T2', 'Segundo produto', 'Anterior', 3, 'G', 4])
        raw = io.BytesIO()
        workbook.save(raw)
        response = self.client.post(f'/api/needs/{self.pid}', files={'file': ('necessidade.xlsx', raw.getvalue())})
        self.assertEqual(response.json()['imported'], 1)
        self.client.post('/api/match', data={'need_id': self.need, 'quote_id': self.quote}).raise_for_status()
        audit = self.client.get(f'/api/audit/{self.pid}').json()
        self.assertEqual(audit[0]['action'], 'VALIDAR_VINCULO')


class Authentication(unittest.TestCase):
    def setUp(self):
        from auth import password_hash
        with compras.db() as c:
            c.execute('DELETE FROM app_sessions WHERE username=?', ('admin',))
            c.execute('INSERT OR REPLACE INTO app_users(username,password_hash,must_change) VALUES (?,?,1)', ('admin', password_hash('123456')))
        self.client = TestClient(compras.app, headers={'X-Sonifarma-Request': '1'})

    def login(self, password='123456'):
        return self.client.post('/api/auth/login', data={'username': 'admin', 'password': password})

    def rotate(self):
        return self.client.post('/api/auth/password', data={'current_password': '123456', 'new_password': 'NovaSenha987!'})

    def test_protected_routes(self):
        for path in ['/api/processes', '/api/data/1', '/api/export/1', '/api/drafts/1/xlsx', '/api/audit/1']:
            self.assertEqual(self.client.get(path).status_code, 401)
        self.assertEqual(self.client.post('/api/processes', data={'name': 'Unauthorized'}).status_code, 401)
        self.assertEqual(self.client.get('/', follow_redirects=False).headers['location'], '/login')

    def test_first_login_rotation_logout(self):
        result = self.login()
        self.assertTrue(result.json()['must_change'])
        self.assertIn('HttpOnly', result.headers['set-cookie'])
        self.assertIn('SameSite=strict', result.headers['set-cookie'])
        self.assertEqual(self.client.get('/api/processes').status_code, 403)
        self.rotate().raise_for_status()
        self.assertEqual(self.client.get('/api/processes').status_code, 401)
        self.assertEqual(self.login().status_code, 401)
        self.login('NovaSenha987!').raise_for_status()
        self.assertEqual(self.client.get('/api/processes').status_code, 200)
        token = self.client.cookies.get('sonifarma_compras_session')
        self.client.post('/api/auth/logout').raise_for_status()
        other = TestClient(compras.app)
        other.cookies.set('sonifarma_compras_session', token)
        self.assertEqual(other.get('/api/processes').status_code, 401)

    def test_csrf_and_lockout(self):
        anonymous = TestClient(compras.app)
        self.assertEqual(anonymous.post('/api/auth/login', data={'username': 'admin', 'password': '123456'}).status_code, 403)
        self.assertEqual(self.client.post('/api/auth/login', data={'username': 'admin', 'password': '123456'}, headers={'Origin': 'https://other.example'}).status_code, 403)
        for _ in range(5):
            self.assertEqual(self.login('wrong').status_code, 401)
        self.assertEqual(self.login().status_code, 429)

    def test_password_validation_and_session_expiry(self):
        self.login().raise_for_status()
        self.assertEqual(self.client.post('/api/auth/password', data={'current_password': 'wrong', 'new_password': 'NovaSenha987!'}).status_code, 400)
        self.assertEqual(self.client.post('/api/auth/password', data={'current_password': '123456', 'new_password': 'short'}).status_code, 400)
        with compras.db() as c:
            stored = c.execute('SELECT password_hash FROM app_users WHERE username=?', ('admin',)).fetchone()[0]
            self.assertNotEqual(stored, '123456')
            c.execute('UPDATE app_sessions SET expires=0 WHERE username=?', ('admin',))
        self.assertFalse(self.client.get('/api/auth/status').json()['authenticated'])
        self.assertEqual(self.client.get('/api/processes').status_code, 401)


if __name__ == '__main__':
    unittest.main()
