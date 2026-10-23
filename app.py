import os,re,sqlite3,unicodedata,math,io
from pathlib import Path
from datetime import datetime
from fastapi import FastAPI,UploadFile,File,Form,HTTPException
from fastapi.responses import HTMLResponse,StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
import fitz
from difflib import SequenceMatcher
from procurement import evaluate, commercial, unit, parse_date, today

ROOT=Path(__file__).parent; DB=Path(os.environ.get('SONIFARMA_COMPRAS_DB', ROOT/'compras.sqlite3'))
app=FastAPI(docs_url=None,redoc_url=None,openapi_url=None,title='Sonifarma | Inteligência de Compras V1.10')
app.mount('/static',StaticFiles(directory=ROOT/'static'),name='static')

def db():
 c=sqlite3.connect(DB);c.row_factory=sqlite3.Row;return c

def init():
 with db() as c:
  c.executescript('''CREATE TABLE IF NOT EXISTS processes(id INTEGER PRIMARY KEY,name TEXT,created TEXT);
 CREATE TABLE IF NOT EXISTS needs(id INTEGER PRIMARY KEY,process_id INTEGER,code TEXT,name TEXT,qty REAL,unit TEXT,last_cost REAL,previous_supplier TEXT);
 CREATE TABLE IF NOT EXISTS quotes(id INTEGER PRIMARY KEY,process_id INTEGER,supplier TEXT,product TEXT,pack TEXT,base_qty REAL,base_unit TEXT,total REAL,unit_cost REAL,expiry TEXT,source TEXT,quote_date TEXT);
 CREATE TABLE IF NOT EXISTS matches(need_id INTEGER,quote_id INTEGER,approved INTEGER DEFAULT 0,PRIMARY KEY(need_id,quote_id));
 CREATE TABLE IF NOT EXISTS choices(need_id INTEGER PRIMARY KEY,quote_id INTEGER,reason TEXT,changed TEXT);
 CREATE TABLE IF NOT EXISTS purchase_audit(id INTEGER PRIMARY KEY,process_id INTEGER,need_id INTEGER,quote_id INTEGER,action TEXT,reason TEXT,created TEXT);''')
init()
with db() as c:
 for table,column,definition in [('needs','last_pack_qty','REAL'),('needs','last_pack_unit','TEXT'),('processes','delivery_date','TEXT')]:
  if column not in {r[1] for r in c.execute('PRAGMA table_info('+table+')')}:
   c.execute('ALTER TABLE '+table+' ADD COLUMN '+column+' '+definition)
 c.execute('CREATE TABLE IF NOT EXISTS supplier_terms(supplier TEXT PRIMARY KEY,minimum_order REAL NOT NULL DEFAULT 700,free_shipping_from REAL NOT NULL DEFAULT 2000,freight REAL,confirmed INTEGER NOT NULL DEFAULT 0)')
from auth import install_auth
install_auth(app, db)

@app.get('/login',response_class=HTMLResponse)
def login_page():return (ROOT/'templates'/'login.html').read_text(encoding='utf-8')

def norm(s):
 s=unicodedata.normalize('NFKD',str(s or '')).encode('ascii','ignore').decode().upper()
 return re.sub(r'[^A-Z0-9]+',' ',s).strip()
def num(v):
 if isinstance(v,(float,int)):return float(v) if math.isfinite(v) else None
 s=str(v or '').replace('R$','').strip().replace(' ','')
 if ',' in s:s=s.replace('.','').replace(',','.')
 try:
  value=float(s)
  return value if math.isfinite(value) else None
 except:return None

def rows_xlsx(raw):
 try:w=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
 except Exception:raise HTTPException(400,'Planilha XLSX inválida ou corrompida')
 try:
  out=[]
  for sheet in w.worksheets:
   rr=[]
   for row in sheet.iter_rows(max_row=1500,max_col=17,values_only=True):
    if any(x is not None for x in row):rr.append(list(row))
   out.append((sheet.title,rr))
  return out
 finally:w.close()

def create_process(name):
 name=name.strip()
 if not name or len(name)>160:raise HTTPException(400,'Informe um nome de processo entre 1 e 160 caracteres')
 with db() as c:
  cur=c.execute('INSERT INTO processes(name,created) VALUES (?,?)',(name,datetime.now().isoformat(timespec='seconds')))
  return cur.lastrowid

def add_need(pid,r):
 code,name,previous,qty,unit,cost=(list(r)+[None]*6)[:6]
 if not name:return
 with db() as c:c.execute('INSERT INTO needs(process_id,code,name,qty,unit,last_cost,previous_supplier) VALUES (?,?,?,?,?,?,?)',(pid,str(code or ''),str(name).strip(),num(qty),str(unit or ''),num(cost),str(previous or '')))

def add_quote(pid,supplier,product,pack,base_qty,base_unit,total,unit_cost,expiry,source,date=''):
 if not product or not supplier:return
 with db() as c:c.execute('INSERT INTO quotes(process_id,supplier,product,pack,base_qty,base_unit,total,unit_cost,expiry,source,quote_date) VALUES (?,?,?,?,?,?,?,?,?,?,?)',(pid,str(supplier),str(product).strip(),str(pack or ''),num(base_qty),str(base_unit or ''),num(total),num(unit_cost),str(expiry or ''),str(source),str(date)))

def seed():
 with db() as c:
  if c.execute('SELECT COUNT(*) FROM processes').fetchone()[0]:return
 source=ROOT/'amostras'/'comparativo_inicial.xlsx'
 if not source.is_file():
  create_process('Novo processo de compras')
  return
 pid=create_process('Outubro 2026 — Dados reais importados')
 data=rows_xlsx(source.read_bytes())
 for sheet,rr in data:
  if sheet=='Necessidade original':
   for r in rr[1:]:add_need(pid,r)
  if sheet=='Todas as cotações':
   for r in rr[1:]:add_quote(pid,r[0],r[1],r[2],r[3],r[4],r[5],r[7],r[8],r[9])
seed()

@app.get('/',response_class=HTMLResponse)
def home():return (ROOT/'templates'/'index.html').read_text(encoding='utf-8')
def require_process(pid):
 with db() as c:
  if not c.execute('SELECT 1 FROM processes WHERE id=?',(pid,)).fetchone():raise HTTPException(404,'Processo não encontrado')

@app.get('/api/processes')
def processes():
 with db() as c:return [dict(r) for r in c.execute('SELECT * FROM processes ORDER BY id DESC')]
@app.post('/api/processes')
def new_process(name:str=Form(...)):return {'id':create_process(name)}
def delivery_date(pid):
 with db() as c:
  r=c.execute('SELECT delivery_date FROM processes WHERE id=?',(pid,)).fetchone()
 return r[0] if r and r[0] else today().isoformat()

def select_champions_automatically(pid,approved_only=False):
 delivery=delivery_date(pid);created=0
 with db() as c:
  needs=[dict(r) for r in c.execute('SELECT * FROM needs WHERE process_id=?',(pid,))]
  quotes=[dict(r) for r in c.execute('SELECT * FROM quotes WHERE process_id=?',(pid,))]
  approved={(r['need_id'],r['quote_id']) for r in c.execute('SELECT * FROM matches WHERE approved=1')}
  for n in needs:
   if c.execute('SELECT 1 FROM choices WHERE need_id=?',(n['id'],)).fetchone():continue
   if not approved_only and c.execute("SELECT 1 FROM purchase_audit WHERE need_id=? AND action='REMOVER'",(n['id'],)).fetchone():continue
   options=[]
   for q in quotes:
    verified=(n['id'],q['id']) in approved
    if not verified and (approved_only or norm(n['name'])!=norm(q['product'])):continue
    assessment=evaluate(n,q,delivery)
    if assessment['eligible']:options.append((not verified,assessment['purchase_total'],assessment['unit_price'],q))
   if not options:continue
   best=min(options,key=lambda item:(item[0],item[1],item[2],item[3]['id']))
   reason='Campeão automático: regras de validade e embalagem atendidas'+(' (conferir)' if best[0] else '')
   now=datetime.now().isoformat(timespec='seconds')
   c.execute('INSERT INTO choices VALUES (?,?,?,?)',(n['id'],best[3]['id'],reason,now))
   c.execute('INSERT INTO purchase_audit(process_id,need_id,quote_id,action,reason,created) VALUES (?,?,?,?,?,?)',(pid,n['id'],best[3]['id'],'CAMPEAO_AUTOMATICO',reason,now))
   created+=1
 return created

@app.get('/api/data/{pid}')
def data(pid:int):
 require_process(pid)
 select_champions_automatically(pid)
 delivery=delivery_date(pid)
 with db() as c:
  needs=[dict(r) for r in c.execute('SELECT * FROM needs WHERE process_id=? ORDER BY name',(pid,))]
  quotes=[dict(r) for r in c.execute('SELECT * FROM quotes WHERE process_id=? ORDER BY supplier,product',(pid,))]
  choices={r['need_id']:dict(r) for r in c.execute('SELECT * FROM choices')}
  approved={(r['need_id'],r['quote_id']) for r in c.execute('SELECT * FROM matches WHERE approved=1')}
 summary={'needs':len(needs),'quotes':len(quotes),'suppliers':len({q['supplier'] for q in quotes}),'selected':0,'pending':0,'champions':0,'selected_total':0,'indicative_saving':0,'comparable_count':0,'review':0}
 for n in needs:
  base,factor=unit(n['unit'],n['name']);n['normalized_unit']=base;n['normalized_qty']=n['qty']*factor if n['qty'] is not None and factor else None
  candidates=[]
  for q in quotes:
   verified=(n['id'],q['id']) in approved
   score=SequenceMatcher(None,norm(n['name']),norm(q['product'])).ratio()
   if verified or score>=.78:
    candidates.append({**q,'score':round(score*100),'verified':verified,'assessment':evaluate(n,q,delivery)})
  candidates.sort(key=lambda q:(not q['assessment']['eligible'],not q['verified'],q['assessment']['purchase_total'] if q['assessment']['purchase_total'] is not None else math.inf,-q['score']))
  comparable=[q for q in candidates if q['assessment']['unit_price'] is not None and (q['verified'] or norm(q['product'])==norm(n['name']))]
  n['lowest_unit']=min(comparable,key=lambda q:q['assessment']['unit_price']) if comparable else None
  n['recommended']=next((q for q in comparable if q['assessment']['eligible']),None)
  n['candidates']=candidates[:25];n['choice']=choices.get(n['id'])
  n['champion']=next(({'id':q['id'],'supplier':q['supplier'],'unit_cost':q['assessment']['unit_price'],'unit':q['assessment']['unit']} for q in candidates if q['verified'] and q['assessment']['eligible']),None)
  summary['champions']+=bool(n['champion'])
  n.update(selected_total=None,saving=None,delta_pct=None,assessment=None,comparison_status='Sem seleção')
  chosen=next((q for q in quotes if n['choice'] and q['id']==n['choice']['quote_id']),None)
  if chosen:
   a=evaluate(n,chosen,delivery);n['assessment']=a;n['selected_total']=a['purchase_total']
   n['comparison_status']='; '.join(a['warnings']) or 'Embalagens inteiras; validade e embalagem compatíveis'
   summary['selected']+=1;summary['selected_total']+=a['purchase_total'] or 0;summary['review']+=not a['eligible']
   if a['unit_price'] is not None and a['quantity'] and n['last_cost'] is not None and n['last_cost']>=0:
    _,factor=unit(n['unit'],n['name']);old=n['last_cost']/factor
    n['saving']=round((old-a['unit_price'])*a['quantity'],2)
    if old>0:n['delta_pct']=round((a['unit_price']/old-1)*100,2)
    summary['indicative_saving']+=n['saving'];summary['comparable_count']+=1
  else:summary['pending']+=1
 summary['selected_total']=round(summary['selected_total'],2);summary['indicative_saving']=round(summary['indicative_saving'],2)
 return {'needs':needs,'quotes':quotes,'summary':summary,'delivery_date':delivery}

@app.post('/api/auto-champions/{pid}')
def auto_champions(pid:int):
 require_process(pid)
 created=select_champions_automatically(pid,approved_only=True)
 return {'selected':created,'pending':0,'warning':'Somente vínculos aprovados com validade superior a um ano, unidades convertíveis e embalagem não maior que a última compra.'}

@app.post('/api/choose')
def choose(need_id:int=Form(...),quote_id:int=Form(...),reason:str=Form('Escolha manual'),override:bool=Form(False)):
 with db() as c:
  n=c.execute('SELECT * FROM needs WHERE id=?',(need_id,)).fetchone()
  q=c.execute('SELECT * FROM quotes WHERE id=?',(quote_id,)).fetchone()
  if not n or not q or n['process_id']!=q['process_id']:raise HTTPException(400,'Produto e cotação incompatíveis')
  assessment=evaluate(dict(n),dict(q),delivery_date(n['process_id']))
  if assessment['warnings']:
   if not override or len(reason.strip())<10:raise HTTPException(409,'; '.join(assessment['warnings'])+'. Confirme a exceção com justificativa de pelo menos 10 caracteres.')
   reason='EXCEÇÃO APROVADA: '+reason.strip()+' | '+ '; '.join(assessment['warnings'])
  c.execute('INSERT OR REPLACE INTO choices VALUES (?,?,?,?)',(need_id,quote_id,reason,datetime.now().isoformat(timespec='seconds')))
  c.execute('INSERT INTO purchase_audit(process_id,need_id,quote_id,action,reason,created) VALUES (?,?,?,?,?,?)',(n['process_id'],need_id,quote_id,'SELECIONAR',reason,datetime.now().isoformat(timespec='seconds')))
 return {'ok':True}
@app.post('/api/match')
def match(need_id:int=Form(...),quote_id:int=Form(...)):
 with db() as c:
  n=c.execute('SELECT process_id FROM needs WHERE id=?',(need_id,)).fetchone()
  q=c.execute('SELECT process_id FROM quotes WHERE id=?',(quote_id,)).fetchone()
  if not n or not q or n[0]!=q[0]:raise HTTPException(400,'Processos diferentes')
  c.execute('INSERT OR REPLACE INTO matches VALUES (?,?,1)',(need_id,quote_id))
  c.execute('INSERT INTO purchase_audit(process_id,need_id,quote_id,action,reason,created) VALUES (?,?,?,?,?,?)',(n[0],need_id,quote_id,'VALIDAR_VINCULO','Equivalência aprovada pelo comprador',datetime.now().isoformat(timespec='seconds')))
 return {'ok':True}
@app.post('/api/needs/{pid}')
async def upload_needs(pid:int,file:UploadFile=File(...)):
 require_process(pid)
 raw=await file.read()
 if not (file.filename or '').lower().endswith('.xlsx'):raise HTTPException(400,'Necessidade deve ser XLSX')
 try:sheets=rows_xlsx(raw)
 except Exception:raise HTTPException(400,'Planilha XLSX inválida ou corrompida')
 count=0
 for _,rr in sheets:
  for i,r in enumerate(rr):
   if len(r)>1 and norm(r[0])=='CODIGO' and 'PRODUTO' in norm(r[1]):
    for item in rr[i+1:]:
     if item[1]:
      add_need(pid,item);count+=1
      headers=[norm(v) for v in r]
      qi=next((j for j,h in enumerate(headers) if h=='EMBALAGEM ANTERIOR QTD'),None)
      ui=next((j for j,h in enumerate(headers) if h=='EMBALAGEM ANTERIOR UNIDADE'),None)
      if qi is not None and ui is not None and num(item[qi]) is not None and num(item[qi])>0:
       with db() as c:c.execute('UPDATE needs SET last_pack_qty=?,last_pack_unit=? WHERE id=(SELECT MAX(id) FROM needs WHERE process_id=?)',(num(item[qi]),str(item[ui] or ''),pid))
    break
 return {'imported':count}
@app.post('/api/quotes/{pid}')
async def upload_quotes(pid:int,supplier:str=Form(...),file:UploadFile=File(...)):
 require_process(pid)
 if not supplier.strip():raise HTTPException(400,'Informe o fornecedor')
 raw=await file.read();fn=(file.filename or '').lower();count=0
 if fn.endswith('.xlsx'):
  for sheet,rr in rows_xlsx(raw):
   if sheet.lower()=='todas as cotações':
    for r in rr[1:]:
     if r[1]:add_quote(pid,r[0] or supplier,r[1],r[2],r[3],r[4],r[5],r[7],r[8],file.filename);count+=1
   else:
    for i,r in enumerate(rr):
     headers=[norm(v) for v in r]
     if any('PRODUTO' in h or 'DESCRICAO' in h for h in headers):
      def idx(keys):return next((j for j,h in enumerate(headers) if any(k in h for k in keys)),None)
      pi=idx(['PRODUTO','DESCRICAO']);price=idx(['TOTAL','VALOR ITEM']);pack=idx(['EMBALAGEM','UNIDADE']);
      for x in rr[i+1:]:
       if pi is not None and x[pi] and price is not None and num(x[price]) is not None:
        add_quote(pid,supplier,x[pi],x[pack] if pack is not None else '',None,'',x[price],None,'',file.filename);count+=1
      break
 elif fn.endswith('.pdf'):
  try:
   with fitz.open(stream=raw,filetype='pdf') as doc:lines='\n'.join(p.get_text() for p in doc).splitlines()
  except Exception:raise HTTPException(400,'Documento PDF inválido ou corrompido')
  # Conservative extraction: only well-structured lines; avoid fabricated associations.
  pattern=re.compile(r'^\s*(.+?)\s{2,}(\d+[\.,]\d{2})\s{2,}(\d+[\.,]\d{2})\s*$')
  for line in lines:
   m=pattern.match(line)
   if m and len(m[1])>4:
    add_quote(pid,supplier,m[1],'',None,'',num(m[3]),None,'',file.filename);count+=1
 else:raise HTTPException(400,'Formato aceito: PDF ou XLSX')
 return {'imported':count,'warning':'Extração de PDF é preliminar. Confira os itens antes de aprovar. PDFs com layouts complexos podem não ser reconhecidos.'}
@app.get('/api/export/{pid}')
def export(pid:int):
 require_process(pid)
 return drafts_pdf(pid)

@app.post('/api/unchoose')
def unchoose(need_id:int=Form(...)):
 with db() as c:
  n=c.execute('SELECT process_id FROM needs WHERE id=?',(need_id,)).fetchone()
  if not n:raise HTTPException(404,'Insumo não encontrado')
  c.execute('DELETE FROM choices WHERE need_id=?',(need_id,))
  c.execute('INSERT INTO purchase_audit(process_id,need_id,action,created) VALUES (?,?,?,?)',(n[0],need_id,'REMOVER',datetime.now().isoformat(timespec='seconds')))
 return {'ok':True}

def draft_rows(pid):
 with db() as c:
  return [dict(r) for r in c.execute("""SELECT n.id AS need_id,n.code,n.name,n.qty,n.unit,n.last_cost,n.last_pack_qty,n.last_pack_unit,q.id AS quote_id,q.supplier,q.product,q.pack,q.base_unit,q.base_qty,q.total,q.unit_cost,q.expiry,ch.reason,ch.changed
  FROM choices ch JOIN needs n ON n.id=ch.need_id JOIN quotes q ON q.id=ch.quote_id WHERE n.process_id=? ORDER BY q.supplier,n.name""",(pid,))]

def prepare_draft(pid):
 groups={};delivery=delivery_date(pid)
 for r in draft_rows(pid):
  a=evaluate(r,r,delivery);r['assessment']=a;r['estimated_cost']=a['purchase_total']
  r['status']='; '.join(a['warnings']) or 'Embalagens inteiras; validade e embalagem compatíveis'
  r['status']+=f" | {a['packages']} embalagem(ns), {a['purchase_qty']} {a['unit']}, excedente {a['surplus_qty']}" if a['packages'] else ''
  groups.setdefault(r['supplier'],[]).append(r)
 return groups

def supplier_conditions(supplier):
 with db() as c:
  r=c.execute('SELECT * FROM supplier_terms WHERE supplier=?',(supplier,)).fetchone()
 return dict(r) if r else {'supplier':supplier,'minimum_order':700,'free_shipping_from':2000,'freight':None,'confirmed':False}

@app.get('/api/drafts/{pid}')
def drafts(pid:int):
 require_process(pid);groups=prepare_draft(pid)
 suppliers=[]
 for supplier,items in groups.items():
  pending=sum(not r['assessment']['eligible'] for r in items)
  total=round(sum(r['estimated_cost'] or 0 for r in items),2)
  suppliers.append({'supplier':supplier,'items':items,'total_estimated':total,'pending':pending,'commercial':commercial(total,supplier_conditions(supplier),pending)})
 return {'suppliers':suppliers,'suggestions':basket_suggestions(pid,groups)}

def basket_suggestions(pid,groups):
 # One-item simulations; never change choices or add unnecessary stock.
 with db() as c:
  quotes=[dict(r) for r in c.execute('SELECT * FROM quotes WHERE process_id=?',(pid,))]
  approved={(r['need_id'],r['quote_id']) for r in c.execute('SELECT * FROM matches WHERE approved=1')}
 totals={name:round(sum(r['estimated_cost'] or 0 for r in rows),2) for name,rows in groups.items()}
 suggestions=[]
 for source,rows in groups.items():
  if any(not r['assessment']['eligible'] for r in rows):continue
  for row in rows:
   if not row['assessment']['eligible']:continue
   for q in quotes:
    target=q['supplier']
    if target==source or (row['need_id'],q['id']) not in approved:continue
    if any(not r['assessment']['eligible'] for r in groups.get(target,[])):continue
    a=evaluate(row,q,delivery_date(pid))
    if not a['eligible'] or a['surplus_qty']>0:continue
    source_after=round(totals[source]-row['estimated_cost'],2);target_after=round(totals.get(target,0)+a['purchase_total'],2)
    st=supplier_conditions(source);tt=supplier_conditions(target)
    if source_after>0 and source_after<st['minimum_order']:continue
    before=commercial(totals.get(target,0),tt);after=commercial(target_after,tt)
    reaches_min=totals.get(target,0)<tt['minimum_order']<=target_after
    reaches_free=totals.get(target,0)<tt['free_shipping_from']<=target_after
    if not (reaches_min or reaches_free):continue
    source_before=commercial(totals[source],st);source_next=commercial(source_after,st)
    delta=None
    if st['confirmed'] and tt['confirmed'] and all(x['landed_total'] is not None for x in [source_before,after]+([before] if totals.get(target,0)>0 else [])+([source_next] if source_after else [])):
     # An empty basket incurs no freight.
     target_before=before['landed_total'] if totals.get(target,0)>0 else 0
     delta=round((source_next['landed_total'] if source_after else 0)+after['landed_total']-source_before['landed_total']-target_before,2)
    suggestions.append({'need_id':row['need_id'],'product':row['name'],'from':source,'to':target,'quote_id':q['id'],'reaches_minimum':reaches_min,'reaches_free_shipping':reaches_free,'target_subtotal':target_after,'source_subtotal':source_after,'merchandise_delta':round(a['purchase_total']-row['estimated_cost'],2),'landed_delta':delta,'note':'Simulação individual; confira as duas cestas antes de aplicar. Sem compra adicional de estoque.'})
 return sorted(suggestions,key=lambda x:(x['landed_delta'] is None,x['landed_delta'] if x['landed_delta'] is not None else x['merchandise_delta']))[:5]

@app.get('/api/audit/{pid}')
def audit(pid:int):
 with db() as c:return [dict(r) for r in c.execute('SELECT * FROM purchase_audit WHERE process_id=? ORDER BY id DESC LIMIT 300',(pid,))]

@app.get('/api/drafts/{pid}/xlsx')
def drafts_xlsx(pid:int):
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill,Alignment
 from openpyxl.utils import get_column_letter
 w=Workbook();w.remove(w.active)
 groups=prepare_draft(pid)
 if not groups:groups={'Sem escolhas':[]}
 for i,(supplier,rows) in enumerate(groups.items()):
  title=re.sub(r'[\\/*?:\[\]]','',supplier).strip()[:27] or 'Fornecedor'
  title=f'{i+1:02d} {title}'[:31]
  sh=w.create_sheet(title)
  sh.append(['SONIFARMA - RASCUNHO DE COMPRAS | NÃO É PEDIDO ENVIADO'])
  sh.append(['Fornecedor:',supplier]);sh.append(['Processo:',str(pid)])
  sh.append(['Código','Insumo necessário','Qtd. necessária','Unidade','Produto cotado','Embalagem','Preço cotado R$','Preço base R$','Total estimado R$','Conferência','Motivo da escolha'])
  for r in rows:sh.append([r['code'],r['name'],r['qty'],r['unit'],r['product'],r['pack'],r['total'],r['unit_cost'],r['estimated_cost'],r['status'],r['reason']])
  terms=commercial(round(sum(r['estimated_cost'] or 0 for r in rows),2),supplier_conditions(supplier),sum(not r['assessment']['eligible'] for r in rows))
  sh.append(['CONDIÇÕES COMERCIAIS',f"Mínimo: {terms['minimum_order']}; falta: {terms['missing_minimum']}; frete grátis: {terms['free_shipping_from']}; falta: {terms['missing_free_shipping']}; frete aplicado: {terms['freight_applied']}; total com frete: {terms['landed_total']}; condições confirmadas: {terms['confirmed']}; pronto: {terms['ready']}"])
  for row in sh:
   for cell in row:
    if isinstance(cell.value,str):cell.data_type='s'
  sh.freeze_panes='A5';sh.auto_filter.ref=f'A4:K{max(4,sh.max_row)}'
  for cell in sh[4]:cell.fill=PatternFill('solid',fgColor='126B52');cell.font=Font(bold=True,color='FFFFFF');cell.alignment=Alignment(wrap_text=True)
  for col,width in enumerate([16,42,18,14,42,23,20,20,22,52,30],1):sh.column_dimensions[get_column_letter(col)].width=width
  for row in sh.iter_rows(min_row=5):
   for idx in [7,8,9]:row[idx-1].number_format='R$ #,##0.00'
   for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
 buffer=io.BytesIO();w.save(buffer);buffer.seek(0)
 return StreamingResponse(buffer,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="sonifarma_rascunhos_{pid}.xlsx"'})

@app.get('/api/drafts/{pid}/pdf')
def drafts_pdf(pid:int):
 from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle,Spacer,PageBreak,KeepTogether
 from reportlab.lib.styles import getSampleStyleSheet
 from reportlab.lib import colors
 from reportlab.lib.pagesizes import landscape,A4
 from xml.sax.saxutils import escape
 buffer=io.BytesIO();doc=SimpleDocTemplate(buffer,pagesize=landscape(A4),leftMargin=24,rightMargin=24,topMargin=25,bottomMargin=25)
 st=getSampleStyleSheet();story=[];groups=prepare_draft(pid)
 for i,(supplier,rows) in enumerate(groups.items()):
  if i:story.append(PageBreak())
  story.append(Paragraph('SONIFARMA | RASCUNHO DE COMPRAS',st['Title']))
  story.append(Paragraph('Fornecedor: '+escape(supplier)+' | Processo: '+str(pid),st['Heading2']))
  story.append(Paragraph('DOCUMENTO PARA CONFERÊNCIA — NÃO REPRESENTA PEDIDO ENVIADO',st['Normal']))
  story.append(Spacer(1,12))
  data=[['Código','Insumo','Necessidade','Produto cotado','Preço base','Total estimado','Situação']]
  for r in rows:data.append([escape(str(r['code'] or '')),Paragraph(escape(str(r['name'] or '')),st['BodyText']),f"{r['qty'] or '?'} {r['unit'] or ''}",Paragraph(escape(str(r['product'] or '')),st['BodyText']),f"R$ {r['unit_cost']:.4f}" if r['unit_cost'] is not None else '-',f"R$ {r['estimated_cost']:.2f}" if r['estimated_cost'] is not None else '-',Paragraph(escape(r['status']),st['BodyText'])])
  t=Table(data,repeatRows=1,colWidths=[65,155,80,170,75,85,155],hAlign='LEFT');t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#126b52')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#eef5f2')]),('FONTSIZE',(0,0),(-1,-1),8),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),7)]));story.append(t)
  total=sum(r['estimated_cost'] or 0 for r in rows);pending=sum(not r['assessment']['eligible'] for r in rows)
  story.append(Spacer(1,12));story.append(Paragraph(f'Total estimado dos itens comparáveis: R$ {total:.2f} | Itens pendentes de conferência: {pending}',st['Heading3']))
  terms=commercial(total,supplier_conditions(supplier),pending)
  story.append(Paragraph(escape(f"Mínimo R$ {terms['minimum_order']:.2f} (falta R$ {terms['missing_minimum']:.2f}); frete grátis a partir de R$ {terms['free_shipping_from']:.2f} (falta R$ {terms['missing_free_shipping']:.2f}); frete: {terms['freight_applied']}; total com frete: {terms['landed_total']}; condições confirmadas: {terms['confirmed']}"),st['Normal']))
  story.append(Paragraph('Conferir equivalência técnica, unidade base, embalagem, pedido mínimo, frete, tributos, validade e disponibilidade antes de aprovar.',st['Normal']))
 if not story:story.append(Paragraph('Nenhuma escolha registrada neste processo.',st['Normal']))
 doc.build(story);buffer.seek(0)
 return StreamingResponse(buffer,media_type='application/pdf',headers={'Content-Disposition':f'attachment; filename="sonifarma_rascunhos_{pid}.pdf"'})

@app.post('/api/processes/{pid}/delivery')
def set_delivery(pid:int,delivery:str=Form(...)):
 require_process(pid)
 parsed=parse_date(delivery)
 if parsed is None:raise HTTPException(400,'Informe uma data de entrega válida')
 with db() as c:
  c.execute('UPDATE processes SET delivery_date=? WHERE id=?',(parsed.isoformat(),pid))
  c.execute('INSERT INTO purchase_audit(process_id,action,reason,created) VALUES (?,?,?,?)',(pid,'ENTREGA_PREVISTA',parsed.isoformat(),datetime.now().isoformat(timespec='seconds')))
 return {'ok':True}

@app.post('/api/needs/{need_id}/package')
def set_last_package(need_id:int,quantity:float=Form(...,gt=0),package_unit:str=Form(...)):
 if not math.isfinite(quantity):raise HTTPException(400,'Quantidade inválida')
 with db() as c:
  n=c.execute('SELECT * FROM needs WHERE id=?',(need_id,)).fetchone()
  if not n:raise HTTPException(404,'Insumo não encontrado')
  base,_=unit(package_unit,n['name']);need_base,_=unit(n['unit'],n['name'])
  if not base or base!=need_base:raise HTTPException(400,'Unidade da embalagem anterior incompatível com o insumo')
  c.execute('UPDATE needs SET last_pack_qty=?,last_pack_unit=? WHERE id=?',(quantity,package_unit,need_id))
  c.execute('INSERT INTO purchase_audit(process_id,need_id,action,reason,created) VALUES (?,?,?,?,?)',(n['process_id'],need_id,'EMBALAGEM_ANTERIOR',f'{quantity} {package_unit}',datetime.now().isoformat(timespec='seconds')))
 return {'ok':True}

@app.post('/api/quotes/{quote_id}/details')
def set_quote_details(quote_id:int,quantity:float=Form(...,gt=0),package_unit:str=Form(...),price:float=Form(...,ge=0),expiry:str=Form(...)):
 if not all(math.isfinite(v) for v in [quantity,price]):raise HTTPException(400,'Quantidade ou preço inválido')
 parsed=parse_date(expiry)
 if parsed is None:raise HTTPException(400,'Informe uma validade válida')
 with db() as c:
  q=c.execute('SELECT * FROM quotes WHERE id=?',(quote_id,)).fetchone()
  if not q:raise HTTPException(404,'Cotação não encontrada')
  if not unit(package_unit,q['product'])[0]:raise HTTPException(400,'Informe g, kg, mg ou milheiro; unidades de cápsulas também são aceitas')
  c.execute('UPDATE quotes SET base_qty=?,base_unit=?,unit_cost=?,total=?,pack=?,expiry=? WHERE id=?',(quantity,package_unit,price,round(quantity*price,2),f'{quantity} {package_unit}',parsed.isoformat(),quote_id))
  c.execute('INSERT INTO purchase_audit(process_id,quote_id,action,reason,created) VALUES (?,?,?,?,?)',(q['process_id'],quote_id,'DADOS_COTACAO',f'Embalagem {quantity} {package_unit}; validade {parsed.isoformat()}; preço/base {price}',datetime.now().isoformat(timespec='seconds')))
 return {'ok':True}

@app.get('/api/suppliers/{pid}')
def supplier_list(pid:int):
 require_process(pid)
 with db() as c:names=[r[0] for r in c.execute('SELECT DISTINCT supplier FROM quotes WHERE process_id=? ORDER BY supplier',(pid,))]
 return [supplier_conditions(name) for name in names]

@app.post('/api/suppliers/conditions')
def set_supplier_terms(supplier:str=Form(...),minimum_order:float=Form(...,ge=0),free_shipping_from:float=Form(...,ge=0),freight:str=Form(''),confirmed:bool=Form(False)):
 minimum=minimum_order;threshold=free_shipping_from;shipping=num(freight)
 if not supplier.strip() or not all(math.isfinite(v) for v in [minimum,threshold]) or (freight.strip() and (shipping is None or shipping<0)):raise HTTPException(400,'Condições comerciais inválidas')
 with db() as c:
  if not c.execute('SELECT 1 FROM quotes WHERE supplier=?',(supplier,)).fetchone():raise HTTPException(404,'Fornecedor não encontrado')
  c.execute('INSERT OR REPLACE INTO supplier_terms VALUES (?,?,?,?,?)',(supplier,minimum,threshold,shipping,int(confirmed)))
  for r in c.execute('SELECT DISTINCT process_id FROM quotes WHERE supplier=?',(supplier,)).fetchall():
   c.execute('INSERT INTO purchase_audit(process_id,action,reason,created) VALUES (?,?,?,?)',(r[0],'CONDICOES_FORNECEDOR',f'{supplier}: mínimo {minimum}; frete grátis {threshold}; frete {shipping}; confirmado {confirmed}',datetime.now().isoformat(timespec='seconds')))
 return {'ok':True}
